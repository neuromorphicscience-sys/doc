from __future__ import annotations

import hashlib
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from app.models import DocumentStructure, ExtractedDocument, ImageBlock, ParagraphBlock, TableBlock


ROLE_STYLE = {
    "title": ("方正小标宋_GBK", 22, WD_ALIGN_PARAGRAPH.CENTER, False),
    "subtitle": ("方正小标宋_GBK", 16, WD_ALIGN_PARAGRAPH.CENTER, False),
    "recipient": ("仿宋_GB2312", 16, WD_ALIGN_PARAGRAPH.JUSTIFY, False),
    "body_paragraph": ("仿宋_GB2312", 16, WD_ALIGN_PARAGRAPH.JUSTIFY, True),
    "heading_level_1": ("黑体", 16, WD_ALIGN_PARAGRAPH.LEFT, False),
    "heading_level_2": ("楷体_GB2312", 16, WD_ALIGN_PARAGRAPH.LEFT, False),
    "heading_level_3": ("仿宋_GB2312", 16, WD_ALIGN_PARAGRAPH.LEFT, False),
    "heading_level_4": ("仿宋_GB2312", 16, WD_ALIGN_PARAGRAPH.LEFT, False),
    "attachment_note": ("仿宋_GB2312", 16, WD_ALIGN_PARAGRAPH.LEFT, True),
    "attachment_body": ("仿宋_GB2312", 16, WD_ALIGN_PARAGRAPH.JUSTIFY, True),
    "issuer_name": ("仿宋_GB2312", 16, WD_ALIGN_PARAGRAPH.RIGHT, False),
    "document_date": ("仿宋_GB2312", 16, WD_ALIGN_PARAGRAPH.RIGHT, False),
    "note": ("仿宋_GB2312", 16, WD_ALIGN_PARAGRAPH.LEFT, True),
    "other": ("仿宋_GB2312", 16, WD_ALIGN_PARAGRAPH.JUSTIFY, False),
}


def _set_run_font(run, east_asia: str, size: int, ascii_font: str | None = None) -> None:
    run.font.name = ascii_font or east_asia
    run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:ascii"), ascii_font or east_asia)
    fonts.set(qn("w:hAnsi"), ascii_font or east_asia)


def _add_text(paragraph, text: str, east_asia: str, size: int, **formatting) -> None:
    for part in re.split(r"([A-Za-z0-9]+)", text):
        if not part:
            continue
        run = paragraph.add_run(part)
        _set_run_font(run, east_asia, size, "Times New Roman" if re.fullmatch(r"[A-Za-z0-9]+", part) else None)
        run.bold = formatting.get("bold", False)
        run.underline = formatting.get("underline", False)
        run.font.superscript = formatting.get("superscript", False)
        run.font.subscript = formatting.get("subscript", False)


def _set_first_line_chars(paragraph, chars: int = 200) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    indent = p_pr.find(qn("w:ind"))
    if indent is None:
        indent = OxmlElement("w:ind")
        p_pr.append(indent)
    indent.set(qn("w:firstLineChars"), str(chars))


def _add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    left = paragraph.add_run("— ")
    _set_run_font(left, "宋体", 14)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))
    _set_run_font(run, "宋体", 14)
    right = paragraph.add_run(" —")
    _set_run_font(right, "宋体", 14)


def _add_paragraph(document: Document, block: ParagraphBlock, role: str) -> None:
    font_name, size, alignment, first_indent = ROLE_STYLE.get(role, ROLE_STYLE["other"])
    paragraph = document.add_paragraph()
    if block.page_break_before:
        paragraph.paragraph_format.page_break_before = True
    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(28.8)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    if first_indent:
        _set_first_line_chars(paragraph)
    for span in block.spans:
        _add_text(
            paragraph,
            span.text,
            font_name,
            size,
            bold=span.bold if role not in {"title", "heading_level_1", "heading_level_2"} else False,
            underline=span.underline,
            superscript=span.superscript,
            subscript=span.subscript,
        )
    if not block.spans and block.text:
        _add_text(paragraph, block.text, font_name, size)


def _add_table(document: Document, block: TableBlock) -> None:
    columns = max((len(row) for row in block.rows), default=1)
    table = document.add_table(rows=len(block.rows), cols=columns)
    table.style = "Table Grid"
    for row_index, row in enumerate(block.rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = Pt(20)
            _add_text(paragraph, value, "仿宋_GB2312", 14)


def _add_image(document: Document, block: ImageBlock, media_dir: Path) -> None:
    image_path = media_dir / block.filename
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    max_width = Mm(156)
    if block.width_emu and block.width_emu < max_width:
        run.add_picture(str(image_path), width=block.width_emu)
    else:
        run.add_picture(str(image_path), width=max_width)


def generate_docx(
    extracted: ExtractedDocument,
    structure: DocumentStructure,
    media_dir: Path,
    output_path: Path,
) -> dict:
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(37)
    section.bottom_margin = Mm(35)
    section.left_margin = Mm(27)
    section.right_margin = Mm(27)
    section.start_type = WD_SECTION.NEW_PAGE

    if document.paragraphs:
        first = document.paragraphs[0]
        first._element.getparent().remove(first._element)

    block_map = {block.id: block for block in extracted.blocks}
    role_map = {item.source_id: item.role for item in structure.items}
    counts: dict[str, int] = {}
    for block in extracted.blocks:
        role = role_map.get(block.id, "other")
        counts[role] = counts.get(role, 0) + 1
        if isinstance(block, ParagraphBlock):
            _add_paragraph(document, block, role)
        elif isinstance(block, TableBlock):
            _add_table(document, block)
        elif isinstance(block, ImageBlock):
            _add_image(document, block, media_dir)

    _add_page_number(section)
    document.core_properties.title = "公文格式规范版"
    document.core_properties.comments = "由公文格式智能规范化系统按模板重建"
    document.save(output_path)
    return counts


def image_hashes_from_docx(path: Path) -> list[str]:
    document = Document(path)
    hashes = []
    for blip in document.element.body.iter(qn("a:blip")):
        rel_id = blip.get(qn("r:embed"))
        relationship = document.part.rels.get(rel_id) if rel_id else None
        target = getattr(relationship, "target_part", None)
        if target is not None and getattr(target, "content_type", "").startswith("image/"):
            hashes.append(hashlib.sha256(target.blob).hexdigest())
    return sorted(hashes)
