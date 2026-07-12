from __future__ import annotations

import hashlib
import mimetypes
import re
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

from app.models import (
    ExtractedDocument,
    ImageBlock,
    InlineSpan,
    ParagraphBlock,
    TableBlock,
)


def _safe_extension(content_type: str, fallback_name: str) -> str:
    extension = mimetypes.guess_extension(content_type) or Path(fallback_name).suffix
    if not extension or not re.fullmatch(r"\.[A-Za-z0-9]{1,8}", extension):
        return ".bin"
    return extension.lower()


def _paragraph_spans(paragraph: Paragraph) -> list[InlineSpan]:
    spans: list[InlineSpan] = []
    for run in paragraph.runs:
        if not run.text:
            continue
        spans.append(
            InlineSpan(
                text=run.text,
                bold=bool(run.bold),
                underline=bool(run.underline),
                superscript=bool(run.font.superscript),
                subscript=bool(run.font.subscript),
            )
        )
    if not spans and paragraph.text:
        spans.append(InlineSpan(text=paragraph.text))
    return spans


def _has_page_break(paragraph: Paragraph) -> bool:
    if paragraph.paragraph_format.page_break_before:
        return True
    for element in paragraph._p.iter(qn("w:br")):
        if element.get(qn("w:type")) == "page":
            return True
    return False


def _paragraph_images(paragraph: Paragraph) -> list[tuple[str, bytes, str, int | None, int | None]]:
    images: list[tuple[str, bytes, str, int | None, int | None]] = []
    rel_ids: list[str] = []
    for blip in paragraph._p.iter(qn("a:blip")):
        rel_id = blip.get(qn("r:embed"))
        if rel_id:
            rel_ids.append(rel_id)

    extents = list(paragraph._p.iter(qn("wp:extent")))
    for index, rel_id in enumerate(rel_ids):
        relation = paragraph.part.rels.get(rel_id)
        if not relation or not hasattr(relation.target_part, "blob"):
            continue
        part = relation.target_part
        width = height = None
        if index < len(extents):
            width = int(extents[index].get("cx")) if extents[index].get("cx") else None
            height = int(extents[index].get("cy")) if extents[index].get("cy") else None
        images.append((Path(part.partname).name, part.blob, part.content_type, width, height))
    return images


def extract_docx(path: Path, media_dir: Path, original_name: str) -> ExtractedDocument:
    document = Document(path)
    media_dir.mkdir(parents=True, exist_ok=True)
    blocks = []
    warnings: list[str] = []
    paragraph_number = table_number = image_number = order = 0

    def append_images(images: list[tuple[str, bytes, str, int | None, int | None]]) -> None:
        nonlocal image_number, order
        for original_filename, blob, content_type, width, height in images:
            image_number += 1
            order += 1
            digest = hashlib.sha256(blob).hexdigest()
            filename = f"image_{image_number:04d}{_safe_extension(content_type, original_filename)}"
            (media_dir / filename).write_bytes(blob)
            blocks.append(
                ImageBlock(
                    id=f"image_{image_number:04d}",
                    order=order,
                    filename=filename,
                    content_type=content_type,
                    sha256=digest,
                    width_emu=width,
                    height_emu=height,
                )
            )

    if any(header.paragraphs and any(p.text.strip() for p in header.paragraphs) for header in (s.header for s in document.sections)):
        warnings.append("检测到页眉文字，首期不会将其作为正文参与结构识别。")
    if any(footer.paragraphs and any(p.text.strip() for p in footer.paragraphs) for footer in (s.footer for s in document.sections)):
        warnings.append("检测到页脚文字，首期生成器会按模板重新创建页脚。")
    if list(document.element.body.iter(qn("wp:anchor"))):
        warnings.append("检测到浮动图片或对象，首期将按普通内嵌图片处理。")

    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            images = _paragraph_images(item)
            if item.text or not images:
                paragraph_number += 1
                order += 1
                blocks.append(
                    ParagraphBlock(
                        id=f"block_{paragraph_number:04d}",
                        order=order,
                        text=item.text,
                        spans=_paragraph_spans(item),
                        page_break_before=_has_page_break(item),
                    )
                )
            append_images(images)
        elif isinstance(item, Table):
            table_number += 1
            order += 1
            blocks.append(
                TableBlock(
                    id=f"table_{table_number:04d}",
                    order=order,
                    rows=[[cell.text for cell in row.cells] for row in item.rows],
                )
            )
            seen_cells: set[int] = set()
            table_images: list[tuple[str, bytes, str, int | None, int | None]] = []
            for row in item.rows:
                for cell in row.cells:
                    cell_identity = id(cell._tc)
                    if cell_identity in seen_cells:
                        continue
                    seen_cells.add(cell_identity)
                    for paragraph in cell.paragraphs:
                        table_images.extend(_paragraph_images(paragraph))
            append_images(table_images)

    return ExtractedDocument(
        filename=original_name,
        blocks=blocks,
        paragraph_count=paragraph_number,
        table_count=table_number,
        image_count=image_number,
        warnings=warnings,
    )
