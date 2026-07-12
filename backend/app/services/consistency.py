from __future__ import annotations

from pathlib import Path

from docx import Document

from app.models import ExtractedDocument, ImageBlock, ParagraphBlock, TableBlock
from app.services.generator import image_hashes_from_docx


def _output_content(path: Path) -> tuple[list[str], list[list[list[str]]]]:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    return paragraphs, tables


def build_consistency_report(extracted: ExtractedDocument, output_path: Path) -> dict:
    source_paragraphs = [block.text for block in extracted.blocks if isinstance(block, ParagraphBlock) and block.text]
    source_tables = [block.rows for block in extracted.blocks if isinstance(block, TableBlock)]
    source_images = sorted(block.sha256 for block in extracted.blocks if isinstance(block, ImageBlock))
    output_paragraphs, output_tables = _output_content(output_path)
    output_images = image_hashes_from_docx(output_path)

    text_match = source_paragraphs == output_paragraphs
    table_match = source_tables == output_tables
    image_match = source_images == output_images
    return {
        "passed": text_match and table_match and image_match,
        "text": {
            "passed": text_match,
            "source_blocks": len(source_paragraphs),
            "output_blocks": len(output_paragraphs),
        },
        "tables": {
            "passed": table_match,
            "source_count": len(source_tables),
            "output_count": len(output_tables),
        },
        "images": {
            "passed": image_match,
            "source_count": len(source_images),
            "output_count": len(output_images),
        },
    }
